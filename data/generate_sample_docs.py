"""
Generates sample PDF and DOCX documents for testing and UI demonstrations.
"""
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import docx

DATA_DIR = Path(__file__).resolve().parent
DOCS_DIR = DATA_DIR / "sample_documents"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

def generate_sample_pdf(file_path: Path):
    c = canvas.Canvas(str(file_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 750, "Rahul Sharma - Curriculum Vitae")
    c.setFont("Helvetica", 10)
    c.drawString(50, 735, "Email: rahul.sharma@example.edu | GitHub: github.com/rahul-ml-dev")
    c.drawString(50, 720, "Location: Bangalore, India")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 695, "EDUCATION")
    c.setFont("Helvetica", 10)
    c.drawString(50, 680, "Bachelor of Technology in Computer Science and Engineering")
    c.drawString(50, 665, "National Institute of Technology | 2021 - 2025")
    c.drawString(50, 650, "Cumulative Grade Point Average (CGPA): 8.2 / 10.0")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 625, "TECHNICAL SKILLS")
    c.setFont("Helvetica", 10)
    c.drawString(50, 610, "Languages & Frameworks: Python, PyTorch, TensorFlow, SQL, C++, Java, Git, Docker")
    c.drawString(50, 595, "Specializations: Machine Learning, Deep Learning, Natural Language Processing, Computer Vision")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 570, "KEY RESEARCH & ENGINEERING PROJECTS")
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 555, "1. Vision-based Plant Disease Detection using Deep CNNs")
    c.setFont("Helvetica", 9)
    c.drawString(60, 540, "- Built ResNet50-based classification pipeline with 96.4% test accuracy across 38 crop disease classes.")
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 520, "2. Multi-Agent Traffic Signal Optimization System")
    c.setFont("Helvetica", 9)
    c.drawString(60, 505, "- Coordinated decentralized Deep Q-Networks to minimize intersection vehicle waiting latency.")
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 485, "3. End-to-End Generative AI Research Summarizer")
    c.setFont("Helvetica", 9)
    c.drawString(60, 470, "- Automated hierarchical paper extraction with RAG and LLM grounded synthesis.")

    c.save()

def generate_sample_docx(file_path: Path):
    doc = docx.Document()
    doc.add_heading("OFFICIAL ACADEMIC TRANSCRIPT / MARKSHEET", level=1)
    doc.add_paragraph("Student Name: Rahul Sharma")
    doc.add_paragraph("Degree: Bachelor of Technology (BTech) in Computer Science")
    doc.add_paragraph("Graduation Status: Expected June 2025")
    doc.add_paragraph("Cumulative Performance Score: CGPA: 8.2 out of 10.0 (First Class with Distinction)")

    doc.add_heading("Key Semester Coursework Completed", level=2)
    courses = [
        ("Data Structures and Algorithms", "Grade: A / 9.0"),
        ("Machine Learning and Pattern Recognition", "Grade: A+ / 10.0"),
        ("Deep Learning and Neural Networks", "Grade: A / 9.0"),
        ("Database Management Systems", "Grade: A / 9.0"),
        ("Probability and Applied Statistics", "Grade: B+ / 8.0"),
        ("Distributed Computing Systems", "Grade: A / 9.0"),
    ]
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Subject Area"
    hdr_cells[1].text = "Academic Evaluation"
    for course, grade in courses:
        row_cells = table.add_row().cells
        row_cells[0].text = course
        row_cells[1].text = grade

    doc.save(str(file_path))

if __name__ == "__main__":
    generate_sample_pdf(DOCS_DIR / "sample_resume.pdf")
    generate_sample_docx(DOCS_DIR / "sample_transcript.docx")
    print("Sample PDF and DOCX successfully generated.")
